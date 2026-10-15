# -*- coding: utf-8 -*-
"""Bangun Presentasi_Kemitraan_KlinikDent.pptx dari hasil analisis kode."""

from pathlib import Path

from lxml import etree
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
PNG = ROOT / "_slot.png"

NAVY = RGBColor(0x0B, 0x2E, 0x4A)
INK = RGBColor(0x17, 0x30, 0x42)
TEAL = RGBColor(0x0E, 0x7C, 0x86)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
PAPER = RGBColor(0xF4, 0xF7, 0xF8)
MUTED = RGBColor(0x5C, 0x6D, 0x7A)
LINE = RGBColor(0xD5, 0xDE, 0xE6)
MIST = RGBColor(0xE8, 0xEE, 0xF2)
SOFT = RGBColor(0xE7, 0xF4, 0xF5)
AMBER = RGBColor(0x8A, 0x5A, 0x10)
AMBER_BG = RGBColor(0xFB, 0xF4, 0xE6)
SLATE = RGBColor(0x3D, 0x51, 0x60)
SLATE_BG = RGBColor(0xE8, 0xEE, 0xF3)
MIST_TEXT = RGBColor(0xD6, 0xE4, 0xEE)
FOOT = RGBColor(0x8A, 0xA0, 0xB0)

SW = 13.333
SH = 7.5
FONT = "Calibri"


def rgb_of(color):
    return (color[0], color[1], color[2]) if False else None


def make_slot():
    w, h = 1920, 1080
    img = Image.new("RGB", (w, h), (232, 238, 242))
    draw = ImageDraw.Draw(img)
    color = (14, 124, 134)
    length, thick, margin = 78, 7, 42
    points = (
        (margin, margin, 1, 1),
        (w - margin, margin, -1, 1),
        (margin, h - margin, 1, -1),
        (w - margin, h - margin, -1, -1),
    )
    for x, y, sx, sy in points:
        draw.line([(x, y), (x + sx * length, y)], fill=color, width=thick)
        draw.line([(x, y), (x, y + sy * length)], fill=color, width=thick)
    img.save(PNG)


def anchor(tf, value):
    body = tf._txBody.find(qn("a:bodyPr"))
    body.set("anchor", value)


def style_run(run, size, bold, color):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = FONT
    rpr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rpr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rpr, qn(tag))
        el.set("typeface", FONT)


def track(run, hundredths):
    run._r.get_or_add_rPr().set("spc", str(hundredths))


def box(slide, x, y, w, h, lines, align="t"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = Inches(0.02)
    tf.margin_right = Inches(0.02)
    tf.margin_top = Inches(0.0)
    tf.margin_bottom = Inches(0.0)
    anchor(tf, {"t": "t", "ctr": "ctr", "b": "b"}[align])
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = line.get("align", PP_ALIGN.LEFT)
        p.space_before = Pt(line.get("before", 0))
        p.space_after = Pt(line.get("after", 0))
        p.line_spacing = line.get("lsp", 1.05)
        run = p.add_run()
        run.text = line["text"]
        style_run(run, line.get("size", 14), line.get("bold", False), line.get("color", INK))
        if line.get("track"):
            track(run, line["track"])
    return shape


def rect(slide, x, y, w, h, fill, line=None, radius=None, name=None):
    kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    if line is None:
        shape.line.fill.background()
    else:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.0)
    if radius is not None:
        try:
            shape.adjustments[0] = radius
        except Exception:
            pass
    if name:
        shape.name = name
    sppr = shape._element.spPr
    effect = sppr.find(qn("a:effectLst"))
    if effect is not None:
        sppr.remove(effect)
    return shape


def shadow(shape):
    sppr = shape._element.spPr
    effect = etree.SubElement(sppr, qn("a:effectLst"))
    outer = etree.SubElement(effect, qn("a:outerShdw"))
    outer.set("blurRad", "63500")
    outer.set("dist", "25400")
    outer.set("dir", "2700000")
    outer.set("algn", "tl")
    outer.set("rotWithShape", "0")
    color = etree.SubElement(outer, qn("a:srgbClr"))
    color.set("val", "0B2E4A")
    alpha = etree.SubElement(color, qn("a:alpha"))
    alpha.set("val", "12000")


def oval(slide, x, y, s, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(s), Inches(s))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background()
    return shape


def blank(prs):
    layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(layout)
    bg = slide.background.fill
    bg.solid()
    bg.fore_color.rgb = PAPER
    return slide


def footer(slide, number):
    rect(slide, 0.42, 7.08, 12.5, 0.01, LINE)
    box(slide, 0.42, 7.12, 8.2, 0.28, [{"text": "KlinikDent.id   ·   Diskusi kemitraan", "size": 11, "color": MUTED}])
    box(
        slide,
        10.4,
        7.12,
        2.5,
        0.28,
        [{"text": f"{number}  /  7", "size": 11, "color": MUTED, "align": PP_ALIGN.RIGHT}],
    )


def title(slide, text, subtitle=None):
    rect(slide, 0, 0, SW, 0.08, TEAL)
    box(slide, 0.42, 0.22, 12.4, 0.48, [{"text": text, "size": 26, "bold": True, "color": NAVY}])
    if subtitle:
        box(slide, 0.42, 0.72, 12.4, 0.36, [{"text": subtitle, "size": 14, "color": MUTED}])


def notes(slide, blocks):
    tf = slide.notes_slide.notes_text_frame
    tf.clear()
    for i, (text, kind) in enumerate(blocks):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_before = Pt(8 if kind == "h" and i else 0)
        p.space_after = Pt(2 if kind == "h" else 4)
        p.level = 0
        run = p.add_run()
        run.text = text
        style_run(run, 15 if kind == "h" else 13, kind == "h", TEAL if kind == "h" else INK)


def shot(slide, x, y, w, h, code, heading, desc):
    """Area gambar 16:9 yang bisa diganti lewat Ubah Gambar. Teks tetap terpisah."""
    frame = rect(slide, x, y, w, h, MIST, TEAL, radius=0.04, name=f"BINGKAI_{code}")
    frame.line.width = Pt(1.25)
    pad = 0.045
    pic = slide.shapes.add_picture(str(PNG), Inches(x + pad), Inches(y + pad), Inches(w - pad * 2), Inches(h - pad * 2))
    pic.name = f"SCREENSHOT_{code}"
    box(
        slide,
        x,
        y + h + 0.08,
        w,
        0.28,
        [{"text": f"SCREENSHOT {code}  —  {heading}", "size": 13, "bold": True, "color": NAVY}],
    )
    box(slide, x, y + h + 0.36, w, 0.58, [{"text": desc, "size": 12, "color": MUTED, "lsp": 1.05}])


def pill(slide, x, y, w, h, text, fg, bg):
    shape = rect(slide, x, y, w, h, bg, radius=0.5)
    box(
        slide,
        x,
        y,
        w,
        h,
        [{"text": text, "size": 11, "bold": True, "color": fg, "align": PP_ALIGN.CENTER}],
        align="ctr",
    )
    return shape


def build():
    make_slot()
    prs = Presentation()
    prs.slide_width = Inches(SW)
    prs.slide_height = Inches(SH)
    prs.core_properties.title = "KlinikDent.id — Diskusi kemitraan"
    prs.core_properties.subject = "Presentasi kemitraan sistem informasi klinik gigi"
    prs.core_properties.category = "Presentasi bisnis"
    slide_cover(prs)
    slide_modules(prs)
    slide_demo(prs)
    slide_ops(prs)
    slide_integrations(prs)
    slide_next_needs(prs)
    slide_close(prs)
    out = ROOT / "Presentasi_Kemitraan_KlinikDent.pptx"
    tmp = ROOT / "_Presentasi_Kemitraan_KlinikDent.pptx"
    prs.save(tmp)
    audit(prs)
    PNG.unlink(missing_ok=True)
    tmp.replace(out)
    return out


def slide_cover(prs):
    s = blank(prs)
    rect(s, 0, 0, 5.22, SH, NAVY)
    rect(s, 5.22, 0, 0.07, SH, TEAL)
    box(
        s,
        0.46,
        1.72,
        4.4,
        0.32,
        [{"text": "SISTEM INFORMASI KLINIK GIGI", "size": 12, "bold": True, "color": TEAL, "track": 140}],
    )
    box(s, 0.46, 2.12, 4.5, 0.7, [{"text": "KlinikDent.id", "size": 36, "bold": True, "color": WHITE}])
    rect(s, 0.48, 2.95, 1.15, 0.045, TEAL)
    box(
        s,
        0.46,
        3.2,
        4.4,
        1.35,
        [{
            "text": "Sistem untuk mendukung pelayanan, operasional, keuangan, dan farmasi klinik gigi.",
            "size": 16,
            "color": MIST_TEXT,
            "lsp": 1.15,
        }],
    )
    box(s, 0.46, 5.95, 4.4, 0.28, [{"text": "Kerja sama pendampingan", "size": 14, "bold": True, "color": WHITE}])
    box(s, 0.46, 6.28, 4.4, 0.55, [{"text": "Untuk menyesuaikan kebutuhan klinik.", "size": 13, "color": MIST_TEXT, "lsp": 1.05}])
    box(s, 0.46, 6.88, 4.4, 0.28, [{"text": "10 menit, termasuk demo singkat", "size": 13, "color": FOOT}])

    box(
        s,
        5.62,
        0.42,
        7.2,
        0.3,
        [{"text": "EMPAT AREA UTAMA", "size": 12, "bold": True, "color": TEAL, "track": 120}],
    )
    cards = [
        ("01", "Pelayanan pasien dan rekam medis", "Pendaftaran, kunjungan, odontogram, diagnosis, tindakan, dan resep."),
        ("02", "Operasional klinik", "Dokter, jadwal, ruangan, hak akses, dan dukungan multi-klinik."),
        ("03", "Keuangan", "Tagihan, kasir, pembayaran, dan laporan sesuai fitur yang tersedia."),
        ("04", "Farmasi dan fitur pendukung", "Stok obat, pengadaan, serta dukungan yang disesuaikan bersama klinik."),
    ]
    for i, (num, head, body) in enumerate(cards):
        col, row = i % 2, i // 2
        x = 5.62 + col * 3.68
        y = 0.95 + row * 2.95
        card = rect(s, x, y, 3.5, 2.75, WHITE, radius=0.08)
        shadow(card)
        oval(s, x + 0.22, y + 0.22, 0.42, TEAL)
        box(
            s,
            x + 0.22,
            y + 0.22,
            0.42,
            0.42,
            [{"text": num, "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}],
            align="ctr",
        )
        box(s, x + 0.2, y + 0.82, 3.1, 0.7, [{"text": head, "size": 16, "bold": True, "color": NAVY, "lsp": 1.05}])
        box(s, x + 0.2, y + 1.58, 3.1, 0.95, [{"text": body, "size": 13, "color": MUTED, "lsp": 1.08}])

    notes(s, [
        ("TUJUAN", "h"),
        ("Audiens memahami KlinikDent.id sebagai sistem informasi klinik gigi, dan bahwa ada kerja sama pendampingan untuk menyesuaikan kebutuhan klinik.", "p"),
        ("NASKAH", "h"),
        ("Terima kasih atas waktunya. KlinikDent.id membantu klinik gigi menjalankan pelayanan pasien, operasional, keuangan, dan farmasi dalam satu sistem. Kami juga membuka kerja sama pendampingan untuk menyesuaikan sistem dengan kebutuhan klinik. Hari ini saya menunjukkan cakupan yang sudah berjalan, bagian yang bisa dilihat langsung, lalu kita petakan kebutuhan yang paling penting.", "p"),
        ("TRANSISI", "h"),
        ("Kita mulai dari modul yang sudah tersedia.", "p"),
        ("PERTANYAAN", "h"),
        ("Agar demo-nya relevan: bagian mana yang paling sering menghambat hari ini — pendaftaran, rekam medis, kasir, atau stok obat?", "p"),
        ("DEMO", "h"),
        ("Belum membuka aplikasi. Jangan menampilkan WhatsApp sebagai identitas produk.", "p"),
        ("VERIFIKASI", "h"),
        ("Jangan menyebut produk terbaik atau sudah memenuhi semua kebutuhan klinik. Jangan membahas SATUSEHAT atau BPJS. Durasi slide ini sekitar 30 detik.", "p"),
    ])


def slide_modules(prs):
    s = blank(prs)
    title(s, "Dari pelayanan pasien hingga pengelolaan klinik")
    groups = [
        ("Pelayanan", [
            "Pendaftaran dan kunjungan",
            "Rekam medis dan odontogram",
            "Diagnosis, tindakan, dan resep",
            "Rujukan dan instruksi pasien",
        ]),
        ("Operasional", [
            "Dokter, jadwal, dan ruangan",
            "Peran dan hak akses",
            "Multi-klinik dan cabang",
            "Dashboard sesuai peran",
        ]),
        ("Keuangan", [
            "Tagihan dan kasir",
            "Tunai, kartu, transfer, QRIS",
            "Shift kasir dan piutang",
            "Pendapatan, arus kas, laba rugi",
        ]),
        ("Farmasi", [
            "Resep sampai penyerahan",
            "Stok, batch, dan kedaluwarsa",
            "Pemakaian stok secara FEFO",
            "Supplier dan penerimaan barang",
        ]),
    ]
    for i, (head, items) in enumerate(groups):
        x = 0.4 + i * 3.22
        card = rect(s, x, 1.22, 3.06, 4.55, WHITE, radius=0.06)
        shadow(card)
        rect(s, x, 1.22, 3.06, 0.08, TEAL)
        box(s, x + 0.18, 1.48, 2.7, 0.42, [{"text": head, "size": 18, "bold": True, "color": NAVY}])
        lines = []
        for n, item in enumerate(items):
            lines.append({"text": "–  " + item, "size": 14, "color": INK, "before": 14 if n else 2, "lsp": 1.05})
        box(s, x + 0.16, 2.15, 2.74, 3.3, lines)
    bar = rect(s, 0.4, 6.0, 12.52, 0.88, WHITE, radius=0.06)
    shadow(bar)
    box(
        s, 0.62, 6.12, 12.1, 0.64,
        [
            {"text": "Juga tersedia", "size": 12, "bold": True, "color": TEAL},
            {"text": "Antrean dan layar monitor, booking sesi dokter, serta master data klinik. Rinciannya mengikuti peran pengguna.", "size": 14, "color": INK, "before": 2},
        ],
    )
    footer(s, 2)
    notes(s, [
        ("TUJUAN", "h"),
        ("Audiens melihat empat kelompok modul yang sudah punya alur di aplikasi, tanpa merasa semua kebutuhan klinik otomatis tercakup.", "p"),
        ("NASKAH", "h"),
        ("Yang sudah berjalan terbagi empat. Pelayanan mencakup pendaftaran, kunjungan, rekam medis, odontogram dewasa dan sulung, diagnosis, tindakan, resep obat jadi maupun racikan, rujukan, dan instruksi pasien. Operasional mencakup dokter, jadwal, ruangan, hak akses, lebih dari satu klinik, dan cabang. Keuangan mencakup tagihan, kasir, pencatatan pembayaran termasuk QRIS, shift, piutang, serta laporan sampai laba rugi. Farmasi memproses resep sampai obat diserahkan, dengan stok per batch. Obat yang lebih dulu kedaluwarsa dipakai lebih dulu. Pengadaan lewat supplier dan penerimaan barang juga ada.", "p"),
        ("RINCIAN YANG TIDAK MUAT DI SLIDE", "h"),
        ("Pelayanan juga memiliki antrean, loket, layar monitor, booking sesi dokter lewat halaman publik, keluhan, screening, tanda vital, dan pemeriksaan gigi. PDF yang aman disebut: rujukan dan instruksi, dari tab yang tampil. Resume dapat diunduh pasien lewat portal. Surat sakit, finalisasi rekam medis, adendum, dan form pesanan lab di layar dokter ada berkasnya, tetapi belum masuk tab pelayanan — jangan didemokan dan jangan dijanjikan sebelum dicek langsung. Daftar pesanan lab yang sudah ada dapat dilihat di menu farmasi.", "p"),
        ("QRIS di kasir adalah pencatatan metode bayar, bukan pembayaran otomatis ke penyedia QRIS. Pembayaran paket langganan klinik terpisah, lewat Midtrans.", "p"),
        ("TRANSISI", "h"),
        ("Agar tidak berhenti di daftar, saya tunjukkan dua tampilan yang nanti bisa kita lihat langsung.", "p"),
        ("PERTANYAAN", "h"),
        ("Di antara empat kelompok ini, mana yang paling dekat dengan pekerjaan harian Bapak/Ibu?", "p"),
        ("DEMO", "h"),
        ("Belum ada gambar. Siapkan akun yang bisa membuka pendaftaran, pelayanan, kasir, dan farmasi. Akun pemilik klinik paling praktis untuk audiens campuran.", "p"),
        ("VERIFIKASI", "h"),
        ("Jangan menyebut surat sakit, finalisasi, atau pesanan lab dari layar dokter sebagai fitur yang siap ditampilkan. Jangan menyebut SATUSEHAT atau BPJS. Durasi sekitar 1 menit.", "p"),
    ])


def slide_demo(prs):
    s = blank(prs)
    title(s, "Gambaran penggunaan KlinikDent.id", "Contoh alur yang dapat diperlihatkan saat demo.")
    shot(
        s, 0.42, 1.28, 6.15, 3.46, "01", "HALAMAN PENDAFTARAN",
        "Ringkasan kunjungan hari ini dan jalan masuk untuk mencatat pasien sebelum pelayanan.",
    )
    shot(
        s, 6.76, 1.28, 6.15, 3.46, "02", "PELAYANAN DAN ODONTOGRAM",
        "Halaman pelayanan pada tab Odontogram, untuk pencatatan kondisi gigi dewasa dan sulung.",
    )
    bar = rect(s, 0.42, 6.05, 12.5, 0.88, WHITE, radius=0.06)
    shadow(bar)
    box(
        s, 0.62, 6.16, 12.1, 0.66,
        [
            {"text": "Urutan demo", "size": 12, "bold": True, "color": TEAL},
            {"text": "Pendaftaran  →  pelayanan dan odontogram  →  tindakan atau resep bila audiensnya dokter  →  kasir dan farmasi.", "size": 14, "color": INK, "before": 2},
        ],
    )
    footer(s, 3)
    notes(s, [
        ("TUJUAN", "h"),
        ("Audiens melihat dua bukti visual, lalu memahami urutan demo. Gambar di slide masih area kosong yang harus diganti screenshot asli.", "p"),
        ("NASKAH", "h"),
        ("Saya sengaja membatasi dua gambar agar tetap terbaca. Yang kiri adalah halaman Pendaftaran: dari sini pasien masuk dan kunjungan hari ini terlihat. Yang kanan adalah halaman Pelayanan, pada tab Odontogram. Identitas pasien ada di bagian atas, lalu kondisi gigi dewasa dan sulung dicatat pada kunjungan yang sama. Tab lain di halaman itu — keluhan, screening, pemeriksaan, diagnosis, tindakan, resep, rujukan, dan instruksi — kita buka sesuai minat audiens, tidak semuanya.", "p"),
        ("TRANSISI", "h"),
        ("Setelah alur klinis, kita lihat bagaimana kasir dan farmasi mendukung hari kerja klinik.", "p"),
        ("PERTANYAAN", "h"),
        ("Apakah pencatatan gigi hari ini masih di kertas, atau sudah ada sistem lain?", "p"),
        ("DEMO", "h"),
        ("Ganti SCREENSHOT_01 dengan halaman Pendaftaran (/registration). Yang harus terlihat: judul halaman, ringkasan kunjungan hari ini, dan akses pendaftaran baru. Jangan memotong judul.", "p"),
        ("Ganti SCREENSHOT_02 dengan satu kunjungan di halaman Pelayanan, tab Odontogram (/visits/{id}/care?tab=odontogram). Pastikan nama pasien dan bagan gigi ikut terpotret. Saat bicara, buka dulu tab Keluhan sebentar, lalu Odontogram.", "p"),
        ("Jika audiens dokter: lanjut ke tab Tindakan dan Resep pada halaman yang sama. Jika audiens pemilik atau operasional: jangan berlama-lama, lanjut ke kasir.", "p"),
        ("Cadangan: Screenshot 03 adalah odontogram bila gambar 02 belum menunjukkannya. Screenshot 04 adalah tab Tindakan atau Resep. Keduanya tidak dipasang di slide agar gambar tetap besar. Cara mengganti ada di Checklist_Screenshot_KlinikDent.md. Klik kanan objek SCREENSHOT_01 atau SCREENSHOT_02, pilih Ubah Gambar. Jangan menghapus judul di bawahnya.", "p"),
        ("VERIFIKASI", "h"),
        ("Pastikan data demo punya pasien, kunjungan, dan beberapa gigi terisi. Jangan membuka tab yang tidak muncul. Durasi bersama demo sekitar 2 menit.", "p"),
    ])


def slide_ops(prs):
    s = blank(prs)
    title(s, "Mendukung aktivitas klinik sehari-hari")
    areas = [
        ("Dokter, jadwal, dan hak akses", "Profil tenaga medis, jadwal praktik, ruangan, serta peran pengguna."),
        ("Tagihan, kasir, dan laporan", "Tagihan dari pelayanan, pencatatan pembayaran, shift kasir, dan laporan keuangan."),
        ("Resep, stok, dan kedaluwarsa", "Resep hingga penyerahan, stok per batch, dan pemakaian yang mengikuti kedaluwarsa."),
    ]
    for i, (head, body) in enumerate(areas):
        x = 0.4 + i * 4.28
        card = rect(s, x, 0.95, 4.12, 1.48, WHITE, radius=0.08)
        shadow(card)
        rect(s, x, 0.95, 0.08, 1.48, TEAL)
        box(s, x + 0.24, 1.06, 3.72, 0.55, [{"text": head, "size": 14, "bold": True, "color": NAVY, "lsp": 1.02}])
        box(s, x + 0.24, 1.62, 3.72, 0.68, [{"text": body, "size": 12, "color": MUTED, "lsp": 1.05}])
    shot(
        s, 0.55, 2.68, 5.9, 3.32, "05", "TAGIHAN DAN KASIR",
        "Detail tagihan: pasien, item pelayanan, dan pencatatan pembayaran.",
    )
    shot(
        s, 6.88, 2.68, 5.9, 3.32, "06", "STOK FARMASI",
        "Stok obat per batch, tanggal kedaluwarsa, dan jumlah yang tersedia.",
    )
    footer(s, 4)
    notes(s, [
        ("TUJUAN", "h"),
        ("Audiens melihat bahwa sistem ikut dipakai setelah pasien selesai diperiksa: siapa yang bertugas, bagaimana uang dicatat, dan bagaimana obat dikendalikan.", "p"),
        ("NASKAH", "h"),
        ("Tiga aktivitas ini yang paling sering ditanyakan pemilik klinik. Pertama, dokter, jadwal, ruangan, dan hak akses. Staf pendaftaran, dokter, farmasi, dan kasir tidak melihat menu yang sama. Kedua, tagihan. Item pelayanan bisa menjadi tagihan, lalu kasir mencatat pembayaran tunai, kartu, transfer, atau QRIS. QRIS di sini dicatat sebagai metode, belum terhubung otomatis ke penyedia pembayaran. Shift kasir dibuka dan ditutup. Laporan pendapatan, arus kas, dan laba rugi tersedia sebagai laporan. Ketiga, farmasi. Resep yang dikirim dokter diproses sampai diserahkan. Stok menyimpan batch dan tanggal kedaluwarsa. Saat obat dikeluarkan, batch yang lebih dulu kedaluwarsa yang terpakai.", "p"),
        ("TRANSISI", "h"),
        ("Selain operasional harian, ada beberapa fitur pendukung. Kebutuhan yang belum tercakup kita bahas lewat pendampingan.", "p"),
        ("PERTANYAAN", "h"),
        ("Pembayaran pasien dan stok obat saat ini dicatat di mana?", "p"),
        ("DEMO", "h"),
        ("Ganti SCREENSHOT_05 dengan detail satu tagihan (/billing/invoices/{id}). Tonjolkan nama pasien, item, dan pilihan metode bayar. Jangan menjanjikan QRIS otomatis.", "p"),
        ("Ganti SCREENSHOT_06 dengan halaman Stok Obat (/pharmacy/stock). Tonjolkan nama obat, batch, tanggal kedaluwarsa, dan jumlah. Jika ditanya FEFO: jelaskan dengan kalimat sederhana, tanpa istilah sistem.", "p"),
        ("VERIFIKASI", "h"),
        ("Pastikan ada tagihan terbuka dan stok dengan batch. Jangan mengklaim sinkronisasi akuntansi real-time. Durasi sekitar 1 menit.", "p"),
    ])


def slide_integrations(prs):
    s = blank(prs)
    title(s, "Fitur pendukung untuk kebutuhan klinik")
    rows = [
        ("WhatsApp", "Pengingat, promo, template, dan riwayat pesan setelah layanan dikonfigurasi.", "Perlu pengaturan", AMBER, AMBER_BG),
        ("Portal pasien", "Pasien melihat kunjungan dan mengunduh resume.", "Tersedia", TEAL, SOFT),
        ("Audit, arsip, langganan", "Audit log, arsip data praktik, dan paket langganan klinik.", "Tersedia", TEAL, SOFT),
    ]
    y = 1.18
    for name, desc, label, fg, bg in rows:
        card = rect(s, 0.42, y, 12.5, 1.12, WHITE, radius=0.06)
        shadow(card)
        box(s, 0.68, y + 0.2, 8.6, 0.36, [{"text": name, "size": 18, "bold": True, "color": NAVY}])
        box(s, 0.68, y + 0.6, 8.8, 0.36, [{"text": desc, "size": 14, "color": MUTED}])
        pill(s, 10.15, y + 0.38, 2.45, 0.36, label, fg, bg)
        y += 1.26
    band = rect(s, 0.42, 5.1, 12.5, 1.7, NAVY, radius=0.06)
    box(
        s, 0.7, 5.32, 11.95, 1.28,
        [
            {"text": "Kerja sama pendampingan", "size": 14, "bold": True, "color": TEAL, "align": PP_ALIGN.CENTER},
            {
                "text": "KlinikDent.id membuka kerja sama pendampingan untuk menyesuaikan kebutuhan klinik.",
                "size": 18,
                "color": WHITE,
                "align": PP_ALIGN.CENTER,
                "before": 6,
            },
        ],
        align="ctr",
    )
    footer(s, 5)
    notes(s, [
        ("TUJUAN", "h"),
        ("Audiens melihat fitur pendukung yang aman dibahas, lalu mendengar bahwa penyesuaian kebutuhan dilakukan lewat kerja sama pendampingan.", "p"),
        ("NASKAH", "h"),
        ("Ada tiga pendukung yang bisa disebut hari ini. WhatsApp untuk pengingat, promo, template, dan riwayat pesan, setelah layanan pesannya dikonfigurasi. Portal pasien untuk melihat kunjungan dan mengunduh resume. Lalu audit log, arsip data praktik, dan paket langganan klinik. Pembayaran paket langganan lewat Midtrans, terpisah dari kasir pasien. Yang belum tercakup tidak saya janjikan di ruangan ini. KlinikDent.id membuka kerja sama pendampingan untuk menyesuaikan kebutuhan klinik.", "p"),
        ("TRANSISI", "h"),
        ("Dari situ, kita lihat contoh kebutuhan yang bisa dibahas dalam pendampingan. Belum tentu semuanya menjadi prioritas klinik ini.", "p"),
        ("PERTANYAAN", "h"),
        ("Apakah pesan ke pasien atau portal pasien relevan untuk operasional klinik ini?", "p"),
        ("DEMO", "h"),
        ("Jangan membuka menu integrasi. Jangan menyebut atau mendemokan SATUSEHAT dan BPJS. Jika ada yang bertanya, jawab singkat: penyesuaian kebutuhan, termasuk integrasi, dibahas dalam kerja sama pendampingan dan tidak menjadi bagian pertemuan ini. Jika ditanya WhatsApp, siapkan Screenshot 08, bukan sebagai pembuka.", "p"),
        ("VERIFIKASI", "h"),
        ("Pastikan tidak ada slide, catatan yang terbaca, atau demo yang menyinggung SATUSEHAT atau BPJS. Durasi sekitar 1 menit.", "p"),
    ])


def slide_next_needs(prs):
    s = blank(prs)
    title(s, "Pengembangan mengikuti kebutuhan klinik", "Penyesuaian dilakukan melalui kerja sama pendampingan, sesuai prioritas klinik.")
    items = [
        ("01", "Rencana perawatan lintas kunjungan", "Instruksi pasien tercatat pada kunjungan yang sedang berlangsung. Rencana yang mengikat beberapa kunjungan belum tersedia."),
        ("02", "Integrasi katalog KFA", "Kode KFA dapat dicatat manual pada data obat. Katalog KFA yang tersinkron belum tersedia."),
        ("03", "Tiket bantuan platform", "Kanal tiket bantuan di dalam aplikasi belum berjalan."),
    ]
    for i, (num, head, body) in enumerate(items):
        x = 0.4 + i * 4.28
        card = rect(s, x, 1.28, 4.12, 2.42, WHITE, radius=0.08)
        shadow(card)
        oval(s, x + 0.18, y_circle := 1.44, 0.38, TEAL)
        box(
            s, x + 0.18, 1.44, 0.38, 0.38,
            [{"text": num, "size": 11, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}],
            align="ctr",
        )
        box(s, x + 0.66, 1.46, 3.25, 0.62, [{"text": head, "size": 14, "bold": True, "color": NAVY, "lsp": 1.02}])
        box(s, x + 0.18, 2.2, 3.76, 1.3, [{"text": body, "size": 13, "color": MUTED, "lsp": 1.05}])
    band = rect(s, 0.4, 3.9, 12.52, 2.85, WHITE, radius=0.08)
    shadow(band)
    box(s, 0.62, 4.06, 12.0, 0.34, [{"text": "Tiga pertanyaan untuk klinik ini", "size": 15, "bold": True, "color": NAVY}])
    questions = [
        ("1", "Proses apa yang masih dikelola secara manual?"),
        ("2", "Kebutuhan mana yang paling mendesak?"),
        ("3", "Apakah ada aplikasi lain yang perlu diintegrasikan?"),
    ]
    for i, (num, q) in enumerate(questions):
        x = 0.62 + i * 4.1
        oval(s, x, 4.58, 0.36, SOFT)
        box(
            s, x, 4.58, 0.36, 0.36,
            [{"text": num, "size": 13, "bold": True, "color": TEAL, "align": PP_ALIGN.CENTER}],
            align="ctr",
        )
        box(s, x, 5.08, 3.85, 1.35, [{"text": q, "size": 15, "color": INK, "lsp": 1.08}])
    footer(s, 6)
    notes(s, [
        ("TUJUAN", "h"),
        ("Audiens memahami tiga contoh pengembangan, tanpa menganggap ketiganya otomatis menjadi kebutuhan atau janji rilis.", "p"),
        ("NASKAH", "h"),
        ("Tiga hal ini belum tersedia sebagai fitur yang bisa dipakai. Saya sebutkan sebagai contoh di dalam pendampingan, bukan sebagai daftar yang harus disetujui. Rencana perawatan yang berlanjut antar kunjungan belum ada. Kode KFA bisa diketik manual, tetapi katalognya belum tersinkron. Tiket bantuan di dalam aplikasi juga belum berjalan. KlinikDent.id membuka kerja sama pendampingan untuk menyesuaikan kebutuhan klinik. Mana yang dikerjakan, kita tentukan bersama.", "p"),
        ("TRANSISI", "h"),
        ("Supaya pertemuan ini menghasilkan langkah yang konkret, saya usulkan empat tindak lanjut.", "p"),
        ("PERTANYAAN", "h"),
        ("Bacakan tiga pertanyaan di slide, lalu diam sejenak. Catat jawaban dengan bahasa klinik, bukan bahasa sistem. Jangan memaksa ketiga kandidat menjadi prioritas.", "p"),
        ("DEMO", "h"),
        ("Tidak ada demo. Jangan membuka halaman tiket: menu itu masih halaman kosong.", "p"),
        ("VERIFIKASI", "h"),
        ("Jangan mengatakan fitur ini sedang dikerjakan atau punya tanggal rilis. Durasi sekitar 1 menit.", "p"),
    ])


def slide_close(prs):
    s = blank(prs)
    title(s, "Tentukan prioritas bersama")
    steps = [
        ("1", "Identifikasi kebutuhan", "Maksimal tiga kebutuhan prioritas dari proses klinik."),
        ("2", "Tentukan PIC", "Satu penanggung jawab dari pihak klinik untuk tindak lanjut."),
        ("3", "Jadwalkan diskusi", "Diskusi teknis lanjutan, 30 menit."),
        ("4", "Susun kesepakatan", "Ruang lingkup, hasil kerja, biaya, dan estimasi waktu setelah analisis."),
    ]
    for i, (num, head, body) in enumerate(steps):
        x = 0.42 + i * 3.22
        card = rect(s, x, 1.28, 3.06, 3.15, WHITE, radius=0.08)
        shadow(card)
        oval(s, x + 1.28, 1.5, 0.5, TEAL)
        box(
            s, x + 1.28, 1.5, 0.5, 0.5,
            [{"text": num, "size": 16, "bold": True, "color": WHITE, "align": PP_ALIGN.CENTER}],
            align="ctr",
        )
        box(s, x + 0.16, 2.18, 2.74, 0.7, [{"text": head, "size": 15, "bold": True, "color": NAVY, "align": PP_ALIGN.CENTER, "lsp": 1.02}])
        box(s, x + 0.16, 2.9, 2.74, 1.25, [{"text": body, "size": 13, "color": MUTED, "align": PP_ALIGN.CENTER, "lsp": 1.08}])
    band = rect(s, 0.42, 4.68, 12.5, 2.15, NAVY, radius=0.06)
    box(
        s, 0.7, 5.05, 11.95, 1.15,
        [{
            "text": "KlinikDent.id membuka kerja sama pendampingan untuk menyesuaikan kebutuhan klinik, berdasarkan prioritas yang disepakati bersama.",
            "size": 16,
            "color": WHITE,
            "align": PP_ALIGN.CENTER,
            "lsp": 1.15,
        }],
        align="ctr",
    )
    box(
        s, 0.7, 6.25, 11.95, 0.32,
        [{"text": "Pertemuan ini menetapkan prioritas. Persetujuan pengembangan menyusul setelah analisis.", "size": 12, "color": MIST_TEXT, "align": PP_ALIGN.CENTER}],
    )
    footer(s, 7)
    notes(s, [
        ("TUJUAN", "h"),
        ("Pertemuan ditutup dengan empat langkah kerja, bukan dengan permintaan menyetujui seluruh fitur.", "p"),
        ("NASKAH", "h"),
        ("Saya usulkan kita tidak menutup pertemuan dengan daftar fitur. Empat langkahnya sederhana. Pertama, pilih maksimal tiga kebutuhan yang benar-benar prioritas. Kedua, tunjuk satu PIC dari klinik. Ketiga, kita jadwalkan diskusi teknis 30 menit. Keempat, setelah kebutuhan itu dianalisis, baru disusun ruang lingkup, hasil kerja, biaya, dan estimasi waktu. KlinikDent.id membuka kerja sama pendampingan untuk menyesuaikan kebutuhan klinik, berdasarkan prioritas yang disepakati bersama.", "p"),
        ("TRANSISI", "h"),
        ("Tidak ada slide berikutnya. Akhiri dengan meminta nama PIC dan dua atau tiga proses yang paling ingin dibereskan.", "p"),
        ("PERTANYAAN", "h"),
        ("Siapa yang tepat menjadi PIC, dan kapan 30 menit diskusi lanjutan bisa dilakukan?", "p"),
        ("DEMO", "h"),
        ("Tidak ada. Jika masih ada waktu, kembali ke satu layar yang tadi paling relevan bagi audiens, bukan menambah fitur baru.", "p"),
        ("VERIFIKASI", "h"),
        ("Jangan menyebutkan harga, tanggal rilis, atau nama pelanggan. Jangan meminta persetujuan atas semua modul. Durasi sekitar 1 menit.", "p"),
        ("CATATAN PENGGANTIAN GAMBAR", "h"),
        ("Objek gambar bernama SCREENSHOT_01, SCREENSHOT_02, SCREENSHOT_05, dan SCREENSHOT_06. Klik kanan objek itu, pilih Ubah Gambar, lalu pilih tangkapan layar asli. Judul dan deskripsi di bawahnya adalah teks terpisah — jangan dihapus. Jangan mengganti seluruh slide menjadi satu gambar.", "p"),
    ])


def audit(prs):
    sw = prs.slide_width
    sh = prs.slide_height
    print(f"slides={len(prs.slides)} size={Emu(sw).inches:.3f}x{Emu(sh).inches:.3f}")
    for i, slide in enumerate(prs.slides, 1):
        note = slide.notes_slide.notes_text_frame.text or ""
        print(f"slide {i}: shapes={len(slide.shapes)} notes={len(note)}")
        if len(note) < 200:
            raise SystemExit(f"notes terlalu pendek pada slide {i}")
        for shape in slide.shapes:
            if shape.has_text_frame:
                visible = shape.text_frame.text or ""
                for word in ("SATUSEHAT", "BPJS", "SatuSehat"):
                    if word.lower() in visible.lower():
                        raise SystemExit(f"teks slide {i} masih menyebut {word}: {visible[:80]}")
            right = shape.left + shape.width
            bottom = shape.top + shape.height
            if shape.left < -2000 or shape.top < -2000 or right > sw + 20000 or bottom > sh + 20000:
                raise SystemExit(
                    f"keluar slide {i} {shape.name} L{shape.left} T{shape.top} R{right} B{bottom}"
                )


if __name__ == "__main__":
    path = build()
    print(path)
